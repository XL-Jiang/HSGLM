import torch.utils.data as utils
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import numpy as np
import torch
import torch.nn.functional as F
from sklearn.model_selection import StratifiedKFold, train_test_split

def save_kfold_indices_to_excel(fold_indices,
                                labels,
                                save_path="kfold_indices.xlsx"):
    """
    兼容两种模式：
      - 5折CV:  fold_indices = [(train_idx, val_idx, test_idx), ...]
      - LOSO:   fold_indices = [(train_idx, val_idx, test_idx, site_name), ...]
    """
    n_samples = len(labels)
    n_splits = len(fold_indices)

    # 判断是 LOSO 还是普通 kfold
    is_loso = len(fold_indices[0]) == 4

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------- Overview ----------
    ws_summary = wb.create_sheet(title="Overview")
    ws_summary.views.sheetView[0].showGridLines = True
    if is_loso:
        ws_summary.append(["Site", "Train Count", "Val Count", "Test Count", "Total Samples"])
    else:
        ws_summary.append(["Fold", "Train Count", "Val Count", "Test Count", "Total Samples"])

    # ---------- Split Matrix ----------
    ws_matrix = wb.create_sheet(title="Split Matrix")
    ws_matrix.views.sheetView[0].showGridLines = True
    if is_loso:
        ws_matrix.append(["Sample Index", "Label"] +
                         [f"Site {fold_indices[i][3]}" for i in range(n_splits)])
    else:
        ws_matrix.append(["Sample Index", "Label"] +
                         [f"Fold {i+1}" for i in range(n_splits)])

    matrix_data = {i: [i, int(labels[i])] + [""] * n_splits for i in range(n_samples)}

    # ---------- 遍历每一折/每个站点 ----------
    for fold, item in enumerate(fold_indices):
        if is_loso:
            train_idx, val_idx, test_idx, site_name = item
            row_header = f"Site {site_name}"
            sheet_name = f"Site_{site_name}"
        else:
            train_idx, val_idx, test_idx = item
            row_header = f"Fold {fold + 1}"
            sheet_name = f"Fold_{fold + 1}"

        ws_summary.append([
            row_header,
            len(train_idx), len(val_idx), len(test_idx),
            len(train_idx) + len(val_idx) + len(test_idx)
        ])

        for idx in train_idx:
            matrix_data[idx][2 + fold] = "Train"
        for idx in val_idx:
            matrix_data[idx][2 + fold] = "Val"
        for idx in test_idx:
            matrix_data[idx][2 + fold] = "Test"

        ws_fold = wb.create_sheet(title=sheet_name[:31])   # Excel sheet 名最长31字符
        ws_fold.views.sheetView[0].showGridLines = True
        ws_fold.append(["Train Index", "Val Index", "Test Index"])
        max_len = max(len(train_idx), len(val_idx), len(test_idx))
        for r in range(max_len):
            tr = int(train_idx[r]) if r < len(train_idx) else ""
            va = int(val_idx[r]) if r < len(val_idx) else ""
            te = int(test_idx[r]) if r < len(test_idx) else ""
            ws_fold.append([tr, va, te])

    for idx in range(n_samples):
        ws_matrix.append(matrix_data[idx])

    # ---------- 样式（与原函数保持一致）----------
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    train_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    val_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    test_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                val_str = str(cell.value)
                if val_str == "Train":
                    cell.fill = train_fill
                elif val_str == "Val":
                    cell.fill = val_fill
                elif val_str == "Test":
                    cell.fill = test_fill
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(save_path)
    print(f"索引已成功导出并保存至: {save_path}")


def init_stratified_kfold_dataloader(cfg, final_timeseires1, final_timeseires2,
                                     final_timeseires3, non_imaging, labels,
                                     stratified, SEED=42, save_excel=True,
                                     n_splits=5, val_ratio=0.2):
    # ---------- 统一转为 tensor，方便后续处理 ----------
    if not isinstance(labels, torch.Tensor):
        labels = torch.tensor(labels)
    if not isinstance(stratified, torch.Tensor):
        stratified = torch.tensor(stratified)

    # 转为 numpy 整数数组，用于构造联合分层标签
    labels_np = labels.cpu().numpy().astype(int)
    stratified_np = stratified.cpu().numpy().astype(int)

    num_classes = labels_np.max() + 1

    stratified_combined = stratified_np * num_classes + labels_np


    labels_onehot = F.one_hot(labels.to(torch.int64))

    skf = StratifiedKFold(n_splits=n_splits, shuffle=True,
                          random_state=getattr(cfg.training, 'seed', SEED))

    folds_dataloaders = []
    fold_indices_list = []

    for fold, (train_val_index, test_index) in enumerate(
            skf.split(final_timeseires1, stratified_combined)):

        train_index, val_index = train_test_split(
            train_val_index,
            test_size=val_ratio,
            random_state=getattr(cfg.training, 'seed', 42) + fold,
            stratify=stratified_combined[train_val_index]  # 联合分层
        )
        fold_indices_list.append((train_index, val_index, test_index))

        # ===== 数据切分 =====
        ts1_train, ts2_train, ts3_train = final_timeseires1[train_index], final_timeseires2[train_index], final_timeseires3[train_index]
        non_img_train, labels_train = non_imaging[train_index], labels_onehot[train_index]

        ts1_val, ts2_val, ts3_val = final_timeseires1[val_index], final_timeseires2[val_index], final_timeseires3[val_index]
        non_img_val, labels_val = non_imaging[val_index], labels_onehot[val_index]

        ts1_test, ts2_test, ts3_test = final_timeseires1[test_index], final_timeseires2[test_index], final_timeseires3[test_index]
        non_img_test, labels_test = non_imaging[test_index], labels_onehot[test_index]

        # 非成像特征标准化（仅用训练折统计量）
        cont_cols = [0, 2, 3]  # 连续特征：年龄、VIQ、PIQ

        mean = non_img_train[:, cont_cols].mean(dim=0, keepdim=True)
        std = non_img_train[:, cont_cols].std(dim=0, keepdim=True) + 1e-6

        def normalize_nonimg(x):
            x = x.clone()
            x[:, cont_cols] = (x[:, cont_cols] - mean) / std
            return x

        non_img_train = normalize_nonimg(non_img_train)
        non_img_val = normalize_nonimg(non_img_val)
        non_img_test = normalize_nonimg(non_img_test)

        train_dataset = utils.TensorDataset(ts1_train, ts2_train, ts3_train, non_img_train, labels_train)
        val_dataset   = utils.TensorDataset(ts1_val,   ts2_val,   ts3_val,   non_img_val,   labels_val)
        test_dataset  = utils.TensorDataset(ts1_test,  ts2_test,  ts3_test,  non_img_test,  labels_test)

        train_dataloader = utils.DataLoader(train_dataset, batch_size=cfg.dataset.batch_size,
                                            shuffle=True, drop_last=cfg.dataset.drop_last)
        val_dataloader   = utils.DataLoader(val_dataset, batch_size=cfg.dataset.batch_size,
                                            shuffle=False, drop_last=False)
        test_dataloader  = utils.DataLoader(test_dataset, batch_size=cfg.dataset.batch_size,
                                            shuffle=False, drop_last=False)

        folds_dataloaders.append((train_dataloader, val_dataloader, test_dataloader))

    if save_excel:
        save_kfold_indices_to_excel(
            fold_indices=fold_indices_list,
            labels=labels_np,
            save_path=f'./fold_index/seed_{SEED}_kfold_indices.xlsx'
        )
    return folds_dataloaders

def init_loso_dataloader(cfg, final_timeseires1, final_timeseires2,
                         final_timeseires3, non_imaging, labels,
                         site_labels, SEED=42, save_excel=True,
                         val_ratio=0.2):
    """
    留一站点交叉验证（LOSO）。
    每个站点依次作为测试集，其余站点作为训练+验证集。
    返回：[(train_loader, val_loader, test_loader, site_name), ...]
    """
    if not isinstance(labels, torch.Tensor):
        labels = torch.tensor(labels)
    if not isinstance(site_labels, torch.Tensor):
        site_labels = torch.tensor(site_labels)
    labels_np = labels.cpu().numpy().astype(int)
    sites_np = site_labels.cpu().numpy().astype(int)
    unique_sites = np.unique(sites_np)

    labels_onehot = F.one_hot(labels.to(torch.int64))

    folds_dataloaders = []
    fold_indices_list = []

    for site in unique_sites:
        test_index = np.where(sites_np == site)[0]
        train_val_index = np.where(sites_np != site)[0]

        # 在训练+验证集内部按诊断分层划分出验证集
        train_index, val_index = train_test_split(
            train_val_index,
            test_size=val_ratio,
            random_state=SEED,
            stratify=labels_np[train_val_index]
        )
        fold_indices_list.append((train_index, val_index, test_index, site))

        # 数据切分
        ts1_train = final_timeseires1[train_index]
        ts2_train = final_timeseires2[train_index]
        ts3_train = final_timeseires3[train_index]
        non_img_train, labels_train = non_imaging[train_index], labels_onehot[train_index]

        ts1_val = final_timeseires1[val_index]
        ts2_val = final_timeseires2[val_index]
        ts3_val = final_timeseires3[val_index]
        non_img_val, labels_val = non_imaging[val_index], labels_onehot[val_index]

        ts1_test = final_timeseires1[test_index]
        ts2_test = final_timeseires2[test_index]
        ts3_test = final_timeseires3[test_index]
        non_img_test, labels_test = non_imaging[test_index], labels_onehot[test_index]

        # 非成像特征标准化（仅用训练折统计量）
        cont_cols = [0, 2, 3]
        mean = non_img_train[:, cont_cols].mean(dim=0, keepdim=True)
        std = non_img_train[:, cont_cols].std(dim=0, keepdim=True) + 1e-6

        def normalize_nonimg(x):
            x = x.clone()
            x[:, cont_cols] = (x[:, cont_cols] - mean) / std
            return x

        non_img_train = normalize_nonimg(non_img_train)
        non_img_val = normalize_nonimg(non_img_val)
        non_img_test = normalize_nonimg(non_img_test)

        # 构建 Dataset
        train_dataset = utils.TensorDataset(ts1_train, ts2_train, ts3_train, non_img_train, labels_train)
        val_dataset = utils.TensorDataset(ts1_val, ts2_val, ts3_val, non_img_val, labels_val)
        test_dataset = utils.TensorDataset(ts1_test, ts2_test, ts3_test, non_img_test, labels_test)

        # 构建 DataLoader
        train_dataloader = utils.DataLoader(train_dataset, batch_size=cfg.dataset.batch_size,
                                            shuffle=True, drop_last=cfg.dataset.drop_last)
        val_dataloader = utils.DataLoader(val_dataset, batch_size=cfg.dataset.batch_size,
                                          shuffle=False, drop_last=False)
        test_dataloader = utils.DataLoader(test_dataset, batch_size=cfg.dataset.batch_size,
                                           shuffle=False, drop_last=False)
        folds_dataloaders.append((train_dataloader, val_dataloader, test_dataloader, f"Site_{site}"))

    if save_excel:
        save_kfold_indices_to_excel(
            fold_indices=fold_indices_list,
            labels=labels_np,
            save_path=f'./fold_index/seed_{SEED}_loso_indices.xlsx'
        )
    return folds_dataloaders